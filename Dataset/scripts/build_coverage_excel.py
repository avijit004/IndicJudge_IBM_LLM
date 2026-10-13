"""
Builds an Excel coverage report from what's actually on disk under dataset_csv/
(the source of truth) — one sheet with a language x task matrix, one sheet
listing every gap with a reason pulled from data/run_report.txt.
"""
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATASET_CSV = ROOT / "dataset_csv"
RUN_REPORT = ROOT / "data" / "run_report.txt"
OUT_XLSX = ROOT / "dataset_coverage.xlsx"

TASKS = ["qa", "reasoning", "summarization", "translation"]
TASK_LABELS = {
    "qa": "Question Answering",
    "reasoning": "Reasoning",
    "summarization": "Summarization",
    "translation": "Translation Quality",
}

LANGUAGES = [
    ("hi", "Hindi", "High"), ("bn", "Bengali", "High"), ("mr", "Marathi", "High"),
    ("te", "Telugu", "High"), ("ta", "Tamil", "High"), ("gu", "Gujarati", "High"),
    ("ur", "Urdu", "High"), ("kn", "Kannada", "High"), ("ml", "Malayalam", "High"),
    ("pa", "Punjabi", "High"),
    ("or", "Odia", "Mid"), ("as", "Assamese", "Mid"), ("ne", "Nepali", "Mid"),
    ("sd", "Sindhi", "Mid"), ("ks", "Kashmiri", "Mid"),
    ("mai", "Maithili", "Low"), ("bho", "Bhojpuri", "Low"), ("gom", "Konkani", "Low"),
    ("doi", "Dogri", "Low"), ("sa", "Sanskrit", "Low"), ("sat", "Santali", "Low"),
    ("brx", "Bodo", "Low"),
]

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
TIER_FILL = PatternFill("solid", fgColor="E5E7EB")
GREEN_FONT = Font(color="006100")
RED_FONT = Font(color="9C0006", bold=True)
HEADER_FONT = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)


def row_count(task: str, iso: str) -> int:
    path = DATASET_CSV / task / f"{iso}.csv"
    if not path.exists():
        return 0
    return len(pd.read_csv(path, encoding="utf-8-sig"))


def parse_gap_reasons() -> dict[tuple[str, str], str]:
    """(task_label, iso) -> short reason, scraped from run_report.txt GAP lines."""
    if not RUN_REPORT.exists():
        return {}
    text = RUN_REPORT.read_text(encoding="utf-8")
    reasons = {}
    blocks = re.split(r"\n(?==== )", text)
    for block in blocks:
        header = re.match(r"=== (.+?) / (.+?) \((\w+)\) ===", block)
        if not header:
            continue
        task_label, _lang, iso = header.groups()
        gap = re.search(r"=> GAP: 0 rows\. \((.+)\)", block)
        if gap:
            reasons[(task_label, iso)] = gap.group(1)
    return reasons


def main() -> None:
    gap_reasons = parse_gap_reasons()

    counts = {task: {iso: row_count(task, iso) for iso, _, _ in LANGUAGES} for task in TASKS}

    # ---- Sheet 1: matrix ----
    matrix_rows = []
    for iso, name, tier in LANGUAGES:
        row = {"Language": name, "ISO": iso, "Resource Tier": tier}
        for task in TASKS:
            n = counts[task][iso]
            row[TASK_LABELS[task]] = n if n else 0
        matrix_rows.append(row)
    df_matrix = pd.DataFrame(matrix_rows)

    # ---- Sheet 2: gap list ----
    gap_rows = []
    for iso, name, tier in LANGUAGES:
        for task in TASKS:
            if counts[task][iso] == 0:
                reason = gap_reasons.get((TASK_LABELS[task], iso), "no dataset found")
                gap_rows.append({"Language": name, "ISO": iso, "Resource Tier": tier,
                                  "Task": TASK_LABELS[task], "Reason": reason})
    df_gaps = pd.DataFrame(gap_rows)

    total_cells = len(LANGUAGES) * len(TASKS)
    filled_cells = sum(1 for iso, _, _ in LANGUAGES for task in TASKS if counts[task][iso] > 0)

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        df_matrix.to_excel(writer, sheet_name="Coverage Matrix", index=False, startrow=1)
        df_gaps.to_excel(writer, sheet_name="Gaps", index=False)

        wb = writer.book

        # --- Coverage Matrix sheet styling ---
        ws = writer.sheets["Coverage Matrix"]
        ws["A1"] = f"Dataset coverage — {filled_cells}/{total_cells} language x task cells filled (100 rows target each)"
        ws["A1"].font = Font(bold=True, size=12)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3 + len(TASKS))

        header_row = 2
        for col_idx, col_name in enumerate(df_matrix.columns, start=1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

        task_col_start = 4  # columns D.. are task columns (1-indexed: Language, ISO, Tier, then tasks)
        for r, (iso, name, tier) in enumerate(LANGUAGES, start=header_row + 1):
            ws.cell(row=r, column=3).fill = TIER_FILL
            for c_offset, task in enumerate(TASKS):
                col = task_col_start + c_offset
                cell = ws.cell(row=r, column=col)
                n = counts[task][iso]
                if n > 0:
                    cell.value = n
                    cell.fill = GREEN
                    cell.font = GREEN_FONT
                else:
                    cell.value = "MISSING"
                    cell.fill = RED
                    cell.font = RED_FONT
                cell.alignment = Alignment(horizontal="center")

        widths = {"A": 14, "B": 8, "C": 14, "D": 22, "E": 14, "F": 18, "G": 20}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "D3"

        # --- Gaps sheet styling ---
        ws2 = writer.sheets["Gaps"]
        for col_idx, col_name in enumerate(df_gaps.columns, start=1):
            cell = ws2.cell(row=1, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        for r in range(2, len(df_gaps) + 2):
            for c in range(1, len(df_gaps.columns) + 1):
                ws2.cell(row=r, column=c).fill = RED
        col_widths2 = {"A": 14, "B": 8, "C": 14, "D": 22, "E": 90}
        for col, w in col_widths2.items():
            ws2.column_dimensions[col].width = w
        ws2.freeze_panes = "A2"

    print(f"Wrote {OUT_XLSX}")
    print(f"Coverage: {filled_cells}/{total_cells} cells filled, {len(df_gaps)} gaps listed")


if __name__ == "__main__":
    main()
