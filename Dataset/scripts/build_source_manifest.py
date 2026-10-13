"""
Builds a manifest of exactly what was pulled from where: for every filled
language x task cell, the source dataset, its config, its license, and how
it was found (original broad scan vs. targeted follow-up research passes).
Reads provenance straight from data/raw/<task>/<iso>.csv (source_dataset /
source_config columns written by collect.py) -- not hand-maintained, so it
can't drift from what actually ran.
"""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
OUT_XLSX = ROOT / "dataset_sources.xlsx"
OUT_CSV = ROOT / "dataset_sources.csv"

TASK_DIRS = {
    "qa": "question_answering",
    "reasoning": "reasoning",
    "summarization": "summarization",
    "translation": "translation_quality",
}
TASK_LABELS = {
    "qa": "Question Answering", "reasoning": "Reasoning",
    "summarization": "Summarization", "translation": "Translation Quality",
}

LANGUAGES = [
    ("hi", "Hindi", "High"), ("bn", "Bengali", "High"), ("mr", "Marathi", "High"),
    ("te", "Telugu", "High"), ("ta", "Tamil", "High"), ("gu", "Gujarati", "High"),
    ("ur", "Urdu", "High"), ("kn", "Kannada", "High"), ("ml", "Malayalam", "High"),
    ("pa", "Punjabi", "High"),
    ("or", "Odia", "Mid"), ("as", "Assamese", "Mid"), ("ne", "Nepali", "Mid"),
    ("sd", "Sindhi", "Mid"), ("ks", "Kashmiri", "Mid"),
    ("mai", "Maithili", "Low"), ("bho", "Bhojpuri", "Low"), ("gom", "Konkani", "Low"),
    ("doi", "Dogri", "Low"), ("sa", "Sanskrit", "Low"), ("sat", "Santali", "Low"),
    ("brx", "Bodo", "Low"),
]

# License + how-it-was-found, keyed by source_dataset string as written into the CSVs.
SOURCE_INFO = {
    "google/IndicGenBench_xorqa_in": ("MIT (page-reported)", "Original broad scan"),
    "google/IndicGenBench_xquad_in": ("CC-BY-SA-4.0", "Original broad scan"),
    "ai4bharat/IndicQA": ("CC-BY-4.0", "Original broad scan"),
    "google-research-datasets/tydiqa": ("Apache-2.0", "Original broad scan"),
    "google/xquad": ("CC-BY-SA-4.0", "Original broad scan"),
    "Yunika/Nepali-QA": ("CC-BY-SA-4.0", "Targeted research pass 2 (individual HF upload)"),
    "l3cube-pune/indic-nlp (L3Cube-IndicQuest)": ("UNCLEAR -- no LICENSE file found, research-use only pending clarification",
                                                    "Targeted research pass 3 (academic paper search)"),
    "ai4bharat/IndicCOPA": ("CC-BY-4.0", "Original broad scan"),
    "bharatgenai/IndicParam": ("CC-BY-NC-4.0", "Original broad scan"),
    "mrlbenchmarks/global-piqa-nonparallel": ("CC-BY-SA-4.0", "Targeted research pass 3 (academic paper search)"),
    "google/IndicGenBench_crosssum_in": ("CC-BY-NC-SA-4.0 + no-train canary", "Original broad scan"),
    "openlanguagedata/flores_plus": ("CC-BY-SA-4.0 (gated -- terms accepted on user's HF account)", "Original broad scan"),
    "google/IndicGenBench_flores_in": ("CC-BY-SA-4.0", "Original broad scan"),
}

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TIER_FILL = PatternFill("solid", fgColor="E5E7EB")


def main() -> None:
    rows = []
    for task, raw_dir in TASK_DIRS.items():
        for iso, name, tier in LANGUAGES:
            path = RAW / raw_dir / f"{iso}.csv"
            if not path.exists():
                continue
            df = pd.read_csv(path, encoding="utf-8-sig")
            if df.empty:
                continue
            src = str(df["source_dataset"].iloc[0])
            cfg = str(df["source_config"].iloc[0])
            license_str, found_via = SOURCE_INFO.get(src, ("UNKNOWN -- not in SOURCE_INFO map", "?"))
            rows.append({
                "Language": name, "ISO": iso, "Resource Tier": tier,
                "Task": TASK_LABELS[task], "Rows": len(df),
                "Source Dataset": src, "Source Config": cfg,
                "License": license_str, "Found Via": found_via,
            })

    manifest = pd.DataFrame(rows).sort_values(["Task", "Resource Tier", "Language"],
                                               key=lambda s: s if s.name != "Resource Tier"
                                               else s.map({"High": 0, "Mid": 1, "Low": 2}))
    manifest.to_csv(OUT_CSV, index=False, encoding="utf-8-sig")

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        manifest.to_excel(writer, sheet_name="Sources", index=False)
        ws = writer.sheets["Sources"]
        for col_idx in range(1, len(manifest.columns) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        widths = {"A": 12, "B": 6, "C": 13, "D": 20, "E": 7, "F": 40, "G": 16, "H": 48, "I": 42}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"

    print(f"Wrote {len(manifest)} rows to {OUT_XLSX} and {OUT_CSV}")
    unknown = manifest[manifest["License"].str.startswith("UNKNOWN")]
    if len(unknown):
        print(f"WARNING: {len(unknown)} rows have a source not in SOURCE_INFO -- update the map.")
        print(unknown[["Task", "Language", "Source Dataset"]].to_string(index=False))


if __name__ == "__main__":
    main()
