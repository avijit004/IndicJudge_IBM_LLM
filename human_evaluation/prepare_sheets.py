"""Build reproducible, blinded reviewer workbooks without calling a model API."""
import ast
from collections import Counter
import csv
import hashlib
import html
import json
from pathlib import Path
import random

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent
SEED = 20260925
RUBRIC = [
    ("Review independently", "Evaluate each response independently. Do not discuss or view another reviewer's ratings before submitting your own."),
    ("What to review", "Evaluate BOTH reasoning_content and response, matching the Sarvam setup. Read the actual P1–P5 instruction for each item."),
    ("Long text", "Open reading_pack.html in a browser and find the evaluation_id (for example HE0001) to read the complete text comfortably. Workbook text is also preserved in full."),
    ("answer_correct", "True if the candidate's final numerical answer matches the supplied reference; False otherwise. The reference is a number, not a human rating."),
    ("reasoning_quality = 0", "No reasoning provided, or completely nonsensical."),
    ("reasoning_quality = 1", "Mostly incorrect. Substantial mathematical or logical errors throughout."),
    ("reasoning_quality = 2", "Partially correct. Some valid reasoning, but significant errors or gaps."),
    ("reasoning_quality = 3", "Mostly correct. The approach is sound with only minor issues."),
    ("reasoning_quality = 4", "Correct and clear. Valid reasoning with no significant errors."),
    ("reasoning_quality = 5", "Excellent. Perfectly structured, mathematically rigorous, and clearly explained."),
    ("follows_prompt_intent", "True if the response satisfies its assigned prompt; False otherwise. Deliberately incorrect reasoning may follow P2/P4 but must still receive an appropriate mathematical-quality score."),
    ("P1", "Correct reasoning and correct final answer."),
    ("P2", "Subtle logical or mathematical reasoning error, but correct final answer."),
    ("P3", "Correct approach with a small arithmetic error at the end producing an incorrect final answer."),
    ("P4", "Incorrect reasoning and incorrect final answer."),
    ("P5", "Correct final answer with only a brief explanation insufficient to fully justify it."),
    ("justification", "Briefly explain all three ratings. Identify the key correct step, error, or missing requirement. Do not judge writing style, language quality, or formatting."),
    ("status", "Choose Complete when reviewer ID, all three ratings, and justification are filled. Use Needs review for ambiguous questions, questionable references, or interpretation issues. Use Cannot assess for language or other barriers; leave unavailable ratings blank."),
    ("flags / notes", "Record reference-answer concerns, ambiguous candidate answers, or conflicts between internal reasoning and the visible response. Resolve rubric ambiguities consistently during calibration."),
    ("Calibration", "First discuss a small separate practice batch and agree on interpretations, then rate these items independently. Preserve original ratings before consensus."),
    ("Reviewer assignment", "Only review languages you can assess confidently. The coordinator may allocate language tabs to qualified reviewers. Enter a reviewer identifier for each rated item."),
]
HEADERS = ["evaluation_id", "prompt_id", "question", "ground_truth", "generation_prompt",
           "reasoning_content", "response", "reviewer_id", "answer_correct", "reasoning_quality",
           "follows_prompt_intent", "justification", "status", "flags / notes"]


def references(path):
    # Use the same answer parser as the judge without importing or executing it.
    import re
    result = {}
    for q in json.loads(path.read_text(encoding="utf-8")):
        answer = q.get("answer_number")
        if answer is None or str(answer).strip() == "":
            match = re.search(r"####\s*([+-]?(?:[0-9]+|[0-9]{1,3}(?:,[0-9]{3})+)(?:\.[0-9]+)?)\s*\Z", q.get("answer", ""))
            if not match:
                raise ValueError(f"Invalid ground truth in {path}: {q['question_id']}")
            answer = match.group(1)
        key = str(q["question_id"])
        assert key not in result
        result[key] = (q["question"], str(answer).strip().replace(",", ""))
    return result


def sample():
    rng = random.Random(SEED)
    selected = []
    for script in sorted((ROOT / "sarvam_judge").glob("sarvam_judge_*.py")):
        tree = ast.parse(script.read_text())
        cfg = {n.targets[0].id: n.value.value for n in tree.body
               if isinstance(n, ast.Assign) and isinstance(n.value, ast.Constant)}
        language = cfg["LANGUAGE"]
        directory = ROOT / cfg["EXPERIMENT_DIR"]
        truth = references(directory / "questions.json")
        language_rows = []
        for model in ("26b", "31b"):
            source = directory / f"gemma4_{model}_{language}_responses.csv"
            with source.open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            for prompt in ("P1", "P2", "P3", "P4", "P5"):
                group = sorted((r for r in rows if r["prompt_id"] == prompt), key=lambda r: int(r["question_id"]))
                assert len(group) == 50
                for row in rng.sample(group, 2):
                    question, answer = truth[row["question_id"]]
                    assert row["question"] == question
                    language_rows.append(dict(row, language=language, ground_truth=answer,
                                              source_model=f"Gemma-4-{model.upper()}",
                                              source_file=str(source.relative_to(ROOT))))
        rng.shuffle(language_rows)
        selected.extend(language_rows)
    for index, row in enumerate(selected, 1):
        row["evaluation_id"] = f"HE{index:04d}"
    assert len(selected) == 200
    assert set(Counter((r["language"], r["source_model"], r["prompt_id"]) for r in selected).values()) == {2}
    return selected


def put(ws, values):
    """Write literal text; candidate content must never become Excel formulas."""
    ws.append(values)
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            if len(cell.value) > 32767:
                raise ValueError("Text exceeds Excel cell limit")
            cell.data_type = "s"


def style(ws, widths, rating_start=None):
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="203864")
        cell.font = Font(color="FFFFFF", bold=True)
    ws.row_dimensions[1].height = 45
    for cell in ws[1]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if rating_start and cell.column >= rating_start:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
        ws.row_dimensions[row[0].row].height = 105 if rating_start else 45


def validation(ws, column, kind, formula1, formula2=None):
    dv = DataValidation(type=kind, formula1=formula1, formula2=formula2,
                        operator="between" if kind == "whole" else None, allow_blank=True)
    dv.errorTitle = "Invalid rating"
    dv.error = "Select an allowed value; reasoning quality must be an integer from 0 to 5."
    dv.showDropDown = False
    dv.showErrorMessage = True
    dv.errorStyle = "stop"
    ws.add_data_validation(dv)
    dv.add(f"{column}2:{column}{ws.max_row}")


def reviewer_book(rows, name):
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    put(ws, ["Topic", "Instructions"])
    for line in RUBRIC:
        put(ws, list(line))
    style(ws, [30, 115])
    for language in sorted({r["language"] for r in rows}):
        ws = wb.create_sheet(language.title())
        put(ws, HEADERS)
        for r in rows:
            if r["language"] == language:
                put(ws, [r["evaluation_id"], r["prompt_id"], r["question"], r["ground_truth"],
                         r["prompt"], r["reasoning_content"], r["response"]] + [None] * 7)
        style(ws, [16, 12, 48, 16, 48, 65, 60, 18, 18, 20, 25, 55, 20, 45], 8)
        validation(ws, "I", "list", '"True,False"')
        validation(ws, "J", "whole", "0", "5")
        validation(ws, "K", "list", '"True,False"')
        validation(ws, "M", "list", '"Complete,Needs review,Cannot assess"')
    wb.save(OUT / name)


def coordinator_book(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Coordinator instructions"
    put(ws, ["Topic", "Instructions"])
    for line in [
        ("Private", "Do not share this workbook with reviewers before independent ratings are submitted: it reveals source models."),
        ("Sampling", f"Seed {SEED}; 200 items; 2 per language × model × prompt; 20 per language. Rows are shuffled within each language."),
        ("Matching", "Use evaluation_id, never row position, to copy each reviewer's ratings into the Consensus sheet."),
        ("Consensus", "Preserve each independent rating. Discuss disagreements or use a third qualified reviewer, then fill consensus ratings and notes. Leave unresolved ratings blank."),
        ("Comparison", "After consensus, join Sarvam results using language, source_model, question_id, prompt_id from Sample key. Compare only completed, assessable items. No Sarvam scores are included here."),
    ]:
        put(ws, line)
    style(ws, [28, 115])
    ws = wb.create_sheet("Sample key")
    keys = ["evaluation_id", "language", "source_model", "question_id", "prompt_id", "source_file", "ground_truth", "input_sha256"]
    put(ws, keys)
    for r in rows:
        digest = hashlib.sha256((ROOT / r["source_file"]).read_bytes()).hexdigest()
        put(ws, [r.get(k, digest) for k in keys])
    style(ws, [18, 18, 20, 16, 14, 75, 18, 70])
    ws = wb.create_sheet("Consensus")
    fields = ["reviewer_id", "answer_correct", "reasoning_quality", "follows_prompt_intent", "justification", "status"]
    put(ws, ["evaluation_id", "language"] + [f"{prefix}_{f}" for prefix in ("r1", "r2") for f in fields]
        + ["consensus_answer_correct", "consensus_reasoning_quality", "consensus_follows_prompt_intent", "adjudicator_id", "resolution_notes"])
    for r in rows:
        put(ws, [r["evaluation_id"], r["language"]] + [None] * 17)
    style(ws, [18, 18] + [25] * 17, 3)
    for col in ("D", "F", "J", "L", "O", "Q"):
        validation(ws, col, "list", '"True,False"')
    for col in ("E", "K", "P"):
        validation(ws, col, "whole", "0", "5")
    wb.save(OUT / "coordinator.xlsx")


def reading_pack(rows):
    parts = ['<!doctype html><html lang="en"><meta charset="utf-8"><title>Human evaluation reading pack</title>',
             '<style>body{font:18px/1.6 system-ui;max-width:1000px;margin:35px auto;padding:0 20px;color:#172334}article{border-top:3px solid #203864;margin:40px 0}pre{white-space:pre-wrap;overflow-wrap:anywhere;font:inherit;background:#f4f6f9;padding:20px}nav a{margin-right:12px}h3{margin-bottom:4px}</style>',
             '<h1>Human evaluation reading pack</h1><p>200 items. Both internal reasoning and visible response are included in full. Find an evaluation ID with your browser search. Enter ratings in your assigned workbook.</p><nav>']
    parts += [f'<a href="#{r["evaluation_id"]}">{r["evaluation_id"]}</a>' for r in rows]
    parts.append('</nav>')
    for r in rows:
        parts.append(f'<article id="{r["evaluation_id"]}"><h2>{r["evaluation_id"]} · {r["language"].title()} · {r["prompt_id"]}</h2>')
        for label, key in [("Question", "question"), ("Ground-truth answer", "ground_truth"), ("Generation instruction", "prompt"), ("Internal reasoning", "reasoning_content"), ("Visible response", "response")]:
            parts.append(f'<h3>{label}</h3><pre>{html.escape(r[key])}</pre>')
        parts.append('</article>')
    parts.append('</html>')
    (OUT / "reading_pack.html").write_text('\n'.join(parts), encoding="utf-8")


def verify(rows):
    expected = {r["evaluation_id"]: r for r in rows}
    for name in ("reviewer_1.xlsx", "reviewer_2.xlsx"):
        wb = load_workbook(OUT / name)
        count = 0
        for ws in wb.worksheets[1:]:
            assert len(ws.data_validations.dataValidation) == 4
            for values in ws.iter_rows(min_row=2, values_only=True):
                r = expected[values[0]]
                assert values[:7] == (r["evaluation_id"], r["prompt_id"], r["question"], r["ground_truth"], r["prompt"], r["reasoning_content"], r["response"])
                assert all(v is None for v in values[7:])
                count += 1
            assert all(c.data_type != "f" for row in ws for c in row)
        assert count == 200
    wb = load_workbook(OUT / "coordinator.xlsx")
    assert wb["Sample key"].max_row == wb["Consensus"].max_row == 201
    print("Verified: 200 balanced items, two blank reviewer workbooks, full text preserved, rating validation, and coordinator matching keys.")


if __name__ == "__main__":
    targets = [OUT / name for name in ("reviewer_1.xlsx", "reviewer_2.xlsx", "coordinator.xlsx", "reading_pack.html")]
    if any(p.exists() for p in targets):
        raise SystemExit("Output already exists. Refusing to overwrite possible human ratings; use a fresh output directory.")
    rows = sample()
    reviewer_book(rows, "reviewer_1.xlsx")
    reviewer_book(rows, "reviewer_2.xlsx")
    coordinator_book(rows)
    reading_pack(rows)
    verify(rows)
